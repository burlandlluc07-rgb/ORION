"""End-to-end test of the *installed* ORION on a real Windows desktop (GitHub Actions runner).

Launches ORION.exe with its native window, drives it through its local API like a user
would, verifies real effects on the PC (processes, files, schedules) and takes screenshots.
Gemini scenarios run only when GEMINI_API_KEY is set.

Outputs e2e/report.md, e2e/report.json and e2e/*.jpg, and prints the screenshots as
base64 log lines ("E2EIMG|name|chunk") so they can be recovered from the job log.
"""

from __future__ import annotations

import base64
import json
import os
import subprocess
import sys
import threading
import time
from pathlib import Path

import httpx
import psutil
from PIL import ImageGrab

EXE = Path(os.environ.get("ORION_EXE", Path(os.environ["LOCALAPPDATA"]) / "Programs" / "ORION" / "ORION.exe"))
HOME = Path(os.environ["ORION_HOME"])
OUT = Path("e2e")
OUT.mkdir(exist_ok=True)
HAS_KEY = bool(os.environ.get("GEMINI_API_KEY"))
DESKTOP = Path.home() / "Desktop"
BG_WORK = HOME / "bg-work"  # the background agent's only allowed folder in this run

for _stream in (sys.stdout, sys.stderr):  # the runner console is cp1252; keep emoji/quotes from crashing prints
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
    except Exception:
        pass

results: list[dict] = []
shots: list[str] = []


def shot(name: str) -> None:
    time.sleep(0.8)
    try:
        img = ImageGrab.grab(all_screens=True).convert("RGB")
        img.thumbnail((1400, 1400))
        img.save(OUT / f"{name}.jpg", "JPEG", quality=62, optimize=True)
        shots.append(name)
    except Exception as exc:
        print(f"screenshot {name} failed: {exc}")


def record(name: str, ok: bool, detail: str, **extra) -> None:
    results.append({"name": name, "ok": ok, "detail": detail, **extra})
    print(f"{'PASS' if ok else 'FAIL'} | {name} | {detail[:300]}", flush=True)


class Orion:
    def __init__(self) -> None:
        info = json.loads((HOME / "instance.json").read_text())
        self.c = httpx.Client(base_url=f"http://127.0.0.1:{info['port']}", headers={"x-orion-token": info["token"]},
                              timeout=60)
        self.conv: str | None = None
        self.approved: list[str] = []
        self._stop = threading.Event()
        threading.Thread(target=self._auto_approve, daemon=True).start()

    def _auto_approve(self) -> None:
        """Plays the user: approves every risky step (after a screenshot of the first approval card)."""
        first = True
        while not self._stop.is_set():
            try:
                for a in self.c.get("/api/state").json().get("approvals", []):
                    if first:
                        shot("05_approval_card")
                        first = False
                    self.c.post(f"/api/approvals/{a['id']}", json={"decision": "approve"})
                    self.approved.append(a["label"])
            except Exception:
                pass
            time.sleep(0.7)

    def ask(self, text: str, timeout: float = 120, mode: str = "auto") -> dict:
        t0 = time.perf_counter()
        r = self.c.post("/api/chat", json={"text": text, "conv_id": self.conv, "mode": mode}).json()
        self.conv = r["conv_id"]
        deadline = time.time() + timeout
        while time.time() < deadline:
            time.sleep(0.6)
            data = self.c.get(f"/api/conversations/{self.conv}").json()
            msgs = data["messages"]
            running = [t for t in data["tasks"] if t["status"] in ("running", "waiting")]
            if msgs and msgs[-1]["role"] == "assistant" and not running:
                last_user = max(i for i, m in enumerate(msgs) if m["role"] == "user")
                if len(msgs) - 1 > last_user:
                    m = msgs[-1]
                    tools = [s["tool"] + ":" + s["status"] for t in data["tasks"] for s in t["steps"]
                             if t["id"] == m["meta"].get("task_id")]
                    return {"reply": m["content"], "model": m["meta"].get("model"), "tools": tools,
                            "seconds": round(time.perf_counter() - t0, 2)}
        return {"reply": "(timeout)", "model": None, "tools": [], "seconds": timeout}

    def window(self, action: str) -> None:
        try:
            self.c.post(f"/api/window/{action}")
        except Exception:
            pass

    def close(self) -> None:
        self._stop.set()


def procs(name: str) -> int:
    return sum(1 for p in psutil.process_iter(["name"]) if (p.info["name"] or "").lower() == name)


def run_tool(o: "Orion", name: str, timeout: float = 300, **args) -> dict:
    try:
        return o.c.post(f"/api/tools/{name}/run", json={"args": args, "timeout": timeout}, timeout=timeout + 30).json()
    except Exception as exc:
        return {"ok": False, "result": repr(exc)}


def toolkit_checks(o: "Orion") -> None:
    """Every new toolkit, called directly on the real PC (no Gemini quota). Effects are verified on disk/OS."""
    docs = Path.home() / "Documents"
    md = "# Test report\n\nRevenue grew **12%**.\n\n- One\n- Two\n\n| Region | Sales |\n|---|---|\n| North | 120 |"
    r1 = run_tool(o, "create_document", path="documents/ORION e2e.docx", content=md, title="E2E", overwrite=True)
    r2 = run_tool(o, "create_document", path="documents/ORION e2e.pdf", content=md + "\n\nCafé ✓ €", overwrite=True)
    r3 = run_tool(o, "create_spreadsheet", path="documents/ORION e2e.xlsx", overwrite=True, chart="bar",
                  rows=[["Item", "Cost"], ["Rent", "1200"], ["Food", "350"], ["Total", "=SUM(B2:B3)"]])
    r4 = run_tool(o, "create_presentation", path="documents/ORION e2e.pptx", title="ORION", overwrite=True,
                  outline="# Why\n- Fast\n- Free\n# Next\n- Ship it")
    back = run_tool(o, "read_file", path="documents/ORION e2e.xlsx")
    files_ok = all((docs / f"ORION e2e.{x}").exists() for x in ("docx", "pdf", "xlsx", "pptx"))
    record("documents: Word, PDF, Excel (formula + chart), PowerPoint", files_ok and all(r["ok"] for r in (r1, r2, r3, r4))
           and "=SUM(B2:B3)" in str(back.get("result")),
           " | ".join(str(r["result"])[:90] for r in (r1, r2, r3, r4)))

    r = run_tool(o, "combine_pdfs", files=["documents/ORION e2e.pdf", "documents/ORION e2e.pdf"],
                 output="documents/ORION combined.pdf")
    z = run_tool(o, "zip_files", paths=["documents/ORION e2e.docx", "documents/ORION e2e.xlsx"],
                 output="documents/ORION e2e.zip")
    x = run_tool(o, "extract_archive", path="documents/ORION e2e.zip", destination="documents/ORION unzip")
    shot_path = Path.home() / "Pictures" / "orion-e2e.png"
    shot_path.parent.mkdir(exist_ok=True)
    ImageGrab.grab().save(shot_path)
    im = run_tool(o, "edit_images", path=str(shot_path), width=640, convert_to="jpg", quality=70)
    dl = run_tool(o, "download_file", url="https://raw.githubusercontent.com/github/gitignore/main/Python.gitignore",
                  folder="downloads")
    srch = run_tool(o, "search_inside_files", text="Revenue grew", folder="documents")
    record("files: PDF merge, zip/unzip, image convert, download, search inside documents",
           all(t["ok"] for t in (r, z, x, im, dl, srch)) and "ORION e2e" in str(srch["result"]),
           " | ".join(str(t["result"])[:70] for t in (r, z, x, im, dl)))

    b1 = run_tool(o, "browser_open", url="https://example.com")
    b2 = run_tool(o, "browser_click", text="Learn more") if b1.get("ok") else b1
    if not b2.get("ok"):
        b2 = run_tool(o, "browser_click", text="More information")
    time.sleep(1)
    shot("10_browser")
    bt = run_tool(o, "browser_tabs")
    record("browser: opens Edge, reads the page, clicks a link", b1.get("ok") and "Example Domain" in str(b1["result"])
           and b2.get("ok") and "iana" in str(b2["result"]).lower(),
           f"{str(b1['result'])[:120]} → {str(b2.get('result'))[:120]} | tabs={str(bt.get('result'))[:80]}")
    run_tool(o, "browser_close")

    th = run_tool(o, "set_theme", mode="dark")
    time.sleep(2)
    shot("11_dark_mode")
    run_tool(o, "set_theme", mode="light")
    wp_img = Path.home() / "Pictures" / "orion-wallpaper.png"
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (1280, 720), (12, 12, 12))
    ImageDraw.Draw(img).text((560, 340), "ORION", fill=(240, 240, 240))
    img.save(wp_img)
    wp = run_tool(o, "set_wallpaper", image=str(wp_img))
    record("windows: dark/light mode and wallpaper (read back)", th["ok"] and wp["ok"],
           f"{th['result']} | {wp['result']}")

    pp = run_tool(o, "power_plan")
    plans = pp.get("result")
    sw = run_tool(o, "power_plan", plan="high performance")
    back_plan = run_tool(o, "power_plan", plan="balanced")
    record("windows: power plans", pp["ok"] and sw["ok"] and back_plan["ok"], f"{str(plans)[:160]} | {sw['result']}")

    ni = run_tool(o, "network_info")
    pg = run_tool(o, "ping", host="1.1.1.1", count=3)
    sp = run_tool(o, "internet_speed_test", timeout=120)
    record("windows: network info, ping, speed test", ni["ok"] and pg["ok"] and sp["ok"] and "download_mbps" in str(sp["result"]),
           f"{str(ni['result'])[:120]} | {str(pg['result'])[:100]} | {sp['result']}")

    dh = run_tool(o, "drive_health")
    su = run_tool(o, "startup_apps")
    ct = run_tool(o, "clean_temp_files")
    rb = run_tool(o, "empty_recycle_bin")
    record("windows: drive health, startup apps, temp cleanup, recycle bin", all(t["ok"] for t in (dh, su, ct, rb)),
           f"{str(dh['result'])[:140]} | {str(su['result'])[:80]} | {ct['result']} | {rb['result']}")

    wifi = run_tool(o, "wifi")
    bt_ = run_tool(o, "bluetooth")
    bat = run_tool(o, "battery_health")
    honest = all((t["ok"] or any(w in str(t["result"]) for w in ("no ", "No ", "desktop"))) for t in (wifi, bt_, bat))
    record("windows: Wi-Fi / Bluetooth / battery answer honestly (runner has none)", honest,
           f"{wifi['result']} | {bt_['result']} | {bat['result']}")

    subprocess.Popen(["notepad.exe"])
    time.sleep(2)
    aw = run_tool(o, "arrange_window", title="Notepad", position="left", second_title="ORION")
    time.sleep(1)
    shot("12_side_by_side")
    record("windows: snap windows side by side", aw["ok"], str(aw["result"]))
    subprocess.run(["taskkill", "/IM", "notepad.exe", "/F"], capture_output=True)
    o.window("maximize")

    wg = run_tool(o, "find_app_to_install", query="7zip")
    record("windows: winget catalog search", wg["ok"] and "7zip" in str(wg["result"]).lower()
           or "isn't available" in str(wg["result"]), str(wg["result"])[:200])

    td = run_tool(o, "todo", action="add", item="Buy milk", list_name="shopping")
    tl = run_tool(o, "todo", action="list", list_name="shopping")
    cc = run_tool(o, "convert_currency", amount=10, from_currency="USD", to_currency="EUR")
    record("productivity: to-do lists and currency", td["ok"] and "Buy milk" in str(tl["result"]) and cc["ok"],
           f"{td['result']} | {cc['result']}")

    wu = run_tool(o, "check_windows_updates", timeout=280)
    record("windows: Windows Update check", wu["ok"], str(wu["result"])[:200])


def memory_checks(o: "Orion") -> None:
    """The Obsidian vault: created in Documents, permanent notes + short-term file, weekly review."""
    root = Path.home() / "Documents" / "ORION Brain"
    o.c.post("/api/memories", json={"text": "Sarah Lane is Lluc's manager at Acme", "kind": "person",
                                    "subject": "Sarah Lane", "permanent": True})
    o.c.post("/api/memories", json={"text": "Prefers metric units", "kind": "preference", "permanent": True})
    o.c.post("/api/memories", json={"text": "Ordered a pizza for tonight", "permanent": False})
    o.c.post("/api/memories", json={"text": "Lluc's sister Maya starts a new job at Google on Monday", "permanent": False})
    st = o.c.get("/api/memory/vault").json()
    people = root / "People" / "Sarah Lane.md"
    ok = (root / "Home.md").exists() and people.exists() and "Prefers metric units" in (root / "About me.md").read_text("utf-8") \
        and "pizza" in (root / "Short-term" / "This week.md").read_text("utf-8") and st["short"] >= 2
    record("memory vault: Obsidian notes in Documents/ORION Brain", ok,
           f"path={st['path']} permanent={st['permanent']} short={st['short']}")
    if HAS_KEY:
        r = o.c.post("/api/memory/review", timeout=240)
        res = r.json() if r.status_code == 200 else {"error": r.text}
        short = (root / "Short-term" / "This week.md").read_text("utf-8")
        record("memory vault: weekly review with Gemini", r.status_code == 200 and "pizza" not in short,
               json.dumps(res)[:300])


def wake_checks() -> None:
    """Custom wake word on real Windows speech recognition (same cases as scripts/wake_check.py)."""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from wake_check import run_cases

    results = run_cases()
    record("custom wake word (Windows offline speech): hears its phrases, ignores other speech",
           all(r["ok"] for r in results),
           " | ".join(f"{r['case']}:{'ok' if r['ok'] else 'FAIL'} hits={[(h['phrase'], h['confidence']) for h in r['hits']]}"
                      for r in results) + f" | lang={results[0]['ready']} err={results[0]['error']}")


def desktop_probe() -> tuple:
    import ctypes
    from ctypes import wintypes

    pt = wintypes.POINT()
    ctypes.windll.user32.GetCursorPos(ctypes.byref(pt))
    return (pt.x, pt.y), int(ctypes.windll.user32.GetForegroundWindow() or 0)


def background_agent_check(o: "Orion") -> None:
    """The in-process background agent on live Gemini: files, a windowless shell command and an overwrite
    that waits for ORION's approval prompt (answered by the auto-approver), with the desktop left alone."""
    (BG_WORK / "notes.txt").write_text("Buy milk, call the dentist on Friday.", "utf-8")
    before = desktop_probe()
    n_approved = len(o.approved)
    t = o.c.post("/api/background/tasks", json={"prompt": (
        "Read notes.txt and write a one-line summary of it to summary.md. Then run the shell command "
        "`echo orion-bg-ok` and include its output in your answer. Finally overwrite notes.txt with the single "
        "word done.")}).json()
    task, t0 = t, time.time()
    while time.time() - t0 < 300 and task.get("status") not in ("completed", "failed", "cancelled"):
        time.sleep(2)
        task = o.c.get(f"/api/background/tasks/{t['id']}").json()
    after = desktop_probe()
    summary = BG_WORK / "summary.md"
    notes = (BG_WORK / "notes.txt").read_text("utf-8").strip().lower()
    asked = [a for a in o.approved[n_approved:] if a.startswith("Background task")]
    ok = (task.get("status") == "completed" and summary.exists() and notes.rstrip(".") == "done" and bool(asked)
          and before == after)
    record("gemini: background agent task (files, shell, approval)", ok,
           f"{task.get('status')} in {time.time() - t0:.0f}s; summary={'yes' if summary.exists() else 'no'}; "
           f"notes={notes[:30]!r}; approvals={len(asked)}; desktop unchanged={before == after}; "
           f"{(task.get('result') or task.get('error') or '')[:200]}")


def main() -> int:
    HOME.mkdir(parents=True, exist_ok=True)
    settings = {"user_name": "Lluc", "onboarded": True, "voice": {"mode": "classic"}}
    if os.environ.get("ORION_BACKUP_API_KEY"):  # the app takes the key from the environment; Gemini stays first
        settings.update(backup_provider=os.environ.get("ORION_BACKUP_PROVIDER") or "groq",
                        backup_model=os.environ.get("ORION_BACKUP_MODEL") or "")
    (HOME / "settings.json").write_text(json.dumps(settings))
    BG_WORK.mkdir(parents=True, exist_ok=True)
    (HOME / "background-agent.toml").write_text(f'[safety]\nallowed_paths = ["{BG_WORK.as_posix()}"]\n', "utf-8")
    proc = subprocess.Popen([str(EXE)], env=os.environ.copy())
    for _ in range(90):
        if (HOME / "instance.json").exists():
            break
        time.sleep(1)
    else:
        record("launch", False, "ORION.exe never became ready")
        shot("00_launch_failed")
        return finish(proc)
    time.sleep(6)  # window + UI load
    o = Orion()
    o.window("maximize")
    st = o.c.get("/api/state").json()
    record("launch installed ORION.exe", True, f"v{st['version']}, voice={st['voice'].get('state')}, key={'yes' if st['has_key'] else 'no'}")
    shot("01_home")

    # ── Instant (local) commands: real effects on the PC ───────────────────
    r = o.ask("open notepad")
    time.sleep(2.5)
    record("fast path: open notepad", procs("notepad.exe") > 0, f"{r['reply']} ({r['model']}, {r['seconds']}s)", **r)
    shot("02_notepad_opened")
    o.window("show")

    r = o.ask("what time is it")
    record("fast path: time", r["model"] == "local" and r["reply"].startswith("It's"), f"{r['reply']} ({r['seconds']}s)", **r)

    r = o.ask("set a timer for 10 minutes")
    timers = o.c.get("/api/routines").json()
    record("fast path: timer persisted", any("Timer" in t["name"] for t in timers), r["reply"], **r)

    r = o.ask("take a screenshot")
    pics = list((Path.home() / "Pictures" / "ORION").glob("*.png"))
    record("fast path: screenshot saved", bool(pics), r["reply"], **r)

    r = o.ask("volume 30")
    record("fast path: volume (honest result)", r["reply"] == "Volume set to 30%." or "can't reach the speakers" in r["reply"],
           r["reply"], **r)

    r = o.ask("how's my pc doing")
    record("fast path: system status", "CPU" in r["reply"], r["reply"], **r)
    o.window("show")
    shot("03_chat_fast_commands")

    # 2.6: settings from an older version (this one has no settings_rev) switch the prompts off once, and
    # closing an app no longer asks (the auto-approver would record it if it did).
    s = o.c.get("/api/state").json()["settings"]
    record("2.6 defaults: approval prompts off", s.get("approval_mode") == "trust" and s.get("desktop_actions") == "allow",
           f"approval_mode={s.get('approval_mode')}, desktop_actions={s.get('desktop_actions')}")
    n_approved = len(o.approved)
    r = o.ask("close notepad")
    time.sleep(2)
    asked = [a for a in o.approved[n_approved:] if "notepad" in a.lower()]
    record("no prompt: close notepad", procs("notepad.exe") == 0 and not asked,
           f"{r['reply']} · prompts: {asked}", **r)

    # ── System check ───────────────────────────────────────────────────────
    checks = o.c.post("/api/diagnostics", json={"generate": HAS_KEY}).json()
    bad = [c for c in checks if c["ok"] is False]
    record("system check", True, "; ".join(f"{c['name']}={'ok' if c['ok'] else ('n/a' if c['ok'] is None else 'FAIL')}"
                                          for c in checks), checks=checks, failures=len(bad))

    # ── Phone app: pair over the LAN server, then use it as a remote ───────
    try:
        rs = o.c.post("/api/remote", json={"enabled": True, "access": "lan"}).json()
        pair_url = rs["pair_url"]
        base = pair_url.split("/#/")[0].replace(rs["lan_url"].split("//")[1].split(":")[0], "127.0.0.1")
        code = pair_url.rsplit("/", 1)[1]
        ph = httpx.Client(base_url=base, timeout=30)
        dev_tok = ph.post("/api/pair", json={"code": code, "name": "E2E phone"}).json()["token"]
        ph.headers["x-orion-token"] = dev_tok
        st_ok = ph.get("/api/state").status_code == 200
        blocked = ph.patch("/api/settings", json={"user_name": "x"}).status_code == 403
        scr = ph.get("/api/screen?w=800")
        status = ph.post("/api/control/system_status", json={}).json()
        app_page = ph.get("/").status_code == 200 and ph.get("/manifest.webmanifest").status_code == 200
        record("phone app: pair + remote control", st_ok and blocked and scr.headers.get("content-type") == "image/jpeg"
               and status.get("ok") and app_page,
               f"paired; settings blocked={blocked}; screen {scr.headers.get('x-screen-size')}; control ok={status.get('ok')}")
        (OUT / "phone_screen.jpg").write_bytes(scr.content)
        shots.append("phone_screen")
    except Exception as exc:
        record("phone app: pair + remote control", False, repr(exc))

    # ── Proactive mode: notices a finished download on its own ─────────────
    dl = Path.home() / "Downloads"
    dl.mkdir(exist_ok=True)
    time.sleep(25)  # let the first scan record a baseline
    (dl / "orion-e2e-download.zip").write_bytes(b"PK" + b"0" * 4096)
    seen = None
    for _ in range(30):
        time.sleep(3)
        seen = next((a for a in o.c.get("/api/alerts").json() if "orion-e2e-download.zip" in a["message"]), None)
        if seen:
            break
    record("proactive: download finished alert", seen is not None, seen["message"] if seen else "no alert within 90s")

    # ── Custom wake word ───────────────────────────────────────────────────
    try:
        wake_checks()
    except Exception as exc:
        record("custom wake word", False, repr(exc))

    # ── Memory vault ───────────────────────────────────────────────────────
    try:
        memory_checks(o)
    except Exception as exc:
        record("memory vault", False, repr(exc))

    # ── New toolkits, called directly ──────────────────────────────────────
    try:
        toolkit_checks(o)
    except Exception as exc:
        record("toolkit checks", False, repr(exc))
        o.window("maximize")

    # ── Gemini scenarios ───────────────────────────────────────────────────
    if HAS_KEY:
        o.conv = None
        target = DESKTOP / "orion-test.txt"
        r = o.ask("Create a text file on my desktop called orion-test.txt containing a two-line poem about the stars.", 180)
        ok = target.exists() and len(target.read_text("utf-8").strip().splitlines()) >= 2
        record("gemini: create file via tools", ok, f"{r['reply'][:200]} | tools={r['tools']} | {r['seconds']}s",
               file=target.read_text("utf-8") if target.exists() else None, **r)
        o.window("show")
        shot("06_gemini_file")

        r = o.ask("Read that file back to me and tell me how many words it has.", 120)
        record("gemini: follow-up uses earlier action", "read_file:done" in r["tools"] or "search_files:done" in r["tools"],
               f"{r['reply'][:200]} | tools={r['tools']}", **r)

        r = o.ask("undo that last file change", 60)
        record("undo: last file change", not target.exists() or "nothing" in r["reply"].lower(), r["reply"][:200], **r)

        subprocess.Popen(["notepad.exe"])
        time.sleep(2)
        subprocess.run(["powershell", "-NoProfile", "-Command",
                        "Add-Type -AssemblyName System.Windows.Forms; Start-Sleep 1; "
                        "[System.Windows.Forms.SendKeys]::SendWait('ORION vision test 7381')"])
        time.sleep(1)
        r = o.ask("Look at my screen: what number is typed in Notepad?", 120)
        record("gemini: screen vision", "7381" in r["reply"].replace(",", "").replace(".", ""), f"{r['reply'][:200]} | tools={r['tools']}", **r)
        subprocess.run(["taskkill", "/IM", "notepad.exe", "/F"], capture_output=True)

        o.conv = None
        r = o.ask("/morning-briefing", 240)
        record("gemini: morning briefing skill", len(r["reply"]) > 80 and "weather" in " ".join(r["tools"]).lower(),
               f"{r['seconds']}s | tools={r['tools']} | {r['reply'][:300]}", **r)
        o.window("show")
        shot("07_morning_briefing")

        r = o.ask("What's the latest news about NASA? Two sentences.", 120)
        record("gemini: web search answer", "web_search:done" in r["tools"] and len(r["reply"]) > 40
               and not r["reply"].lstrip().startswith("⚠"),
               f"{r['reply'][:250]} | tools={r['tools']}", **r)
        shot("08_web_answer")

        o.conv = None
        sheet = DESKTOP / "orion-budget.xlsx"
        r = o.ask("Make an Excel spreadsheet on my desktop called orion-budget.xlsx with rent 1200 and food 350, "
                  "plus a total row that uses a formula.", 180)
        formula = False
        if sheet.exists():
            from openpyxl import load_workbook

            formula = any(str(c.value).startswith("=") for row in load_workbook(sheet).active.iter_rows() for c in row)
        record("gemini: makes a real Excel file (documents toolkit)", sheet.exists() and formula,
               f"{r['reply'][:200]} | tools={r['tools']}", **r)

        o.conv = None
        r = o.ask("Use the browser to open example.com and tell me the main heading on the page.", 180)
        record("gemini: drives the browser (browser toolkit)", any(t.startswith("browser_") for t in r["tools"])
               and "example domain" in r["reply"].lower(), f"{r['reply'][:200]} | tools={r['tools']}", **r)
        run_tool(o, "browser_close")

        report = DESKTOP / "orion-mission.md"
        m = o.c.post("/api/missions", json={"goal": "Find the current weather in London and the top 3 technology news "
                      "headlines today, then save a short Markdown briefing as orion-mission.md on my desktop."}).json()
        status = "queued"
        t0 = time.time()
        while time.time() - t0 < 300:
            time.sleep(4)
            mm = next(x for x in o.c.get("/api/missions").json() if x["id"] == m["id"])
            status = mm["status"]
            if status in ("done", "failed", "stopped"):
                break
        record("gemini: background mission", status == "done" and report.exists(),
               f"{status} in {time.time() - t0:.0f}s; report={'yes' if report.exists() else 'no'}; "
               f"{(mm.get('summary') or '')[:200]}", report=report.read_text('utf-8')[:1500] if report.exists() else None)
        o.window("show")
        shot("09_after_mission")

        background_agent_check(o)

        q = o.c.get("/api/quota").json()["models"]
        record("free-tier usage", True, ", ".join(f"{m['model']}: {m['rpd_used']}/{m['rpd']} today" for m in q))
    else:
        record("gemini scenarios", True, "skipped — add a GEMINI_API_KEY repository secret to run them")

    o.close()
    return finish(proc)


def finish(proc: subprocess.Popen) -> int:
    try:
        proc.terminate()
        proc.wait(10)
    except Exception:
        proc.kill()
    log = HOME / "logs" / "orion.log"
    tail = log.read_text("utf-8", errors="replace").splitlines()[-60:] if log.exists() else []
    passed = sum(r["ok"] for r in results)
    lines = [f"# ORION Windows end-to-end — {passed}/{len(results)} passed", ""]
    lines += [f"- {'✅' if r['ok'] else '❌'} **{r['name']}** — {r['detail']}" for r in results]
    limits = [ln[:600] for ln in (log.read_text("utf-8", errors="replace").splitlines() if log.exists() else [])
              if "Rate limited on" in ln]
    lines += ["", "## Rate limits hit", "```", *(limits or ["none"]), "```"]
    lines += ["", "## Log tail", "```", *tail, "```"]
    (OUT / "report.md").write_text("\n".join(lines), "utf-8")
    (OUT / "report.json").write_text(json.dumps(results, indent=1, default=str), "utf-8")
    print("\n".join(lines))
    print("E2E-SCREENSHOTS-BEGIN")
    for name in shots:
        data = base64.b64encode((OUT / f"{name}.jpg").read_bytes()).decode()
        for i in range(0, len(data), 3500):
            print(f"E2EIMG|{name}|{data[i:i + 3500]}")
    print("E2E-SCREENSHOTS-END")
    return 0 if all(r["ok"] for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())

