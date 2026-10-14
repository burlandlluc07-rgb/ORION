"""Documents toolkit: real Office files, PDFs, images and archives (round-tripped through their readers)."""

import zipfile

import pytest

import orion.tools.files as files
import orion.tools.platform as platform
from orion.tools import media, office
from orion.tools.registry import ToolError


@pytest.fixture
def home(tmp_path, monkeypatch):
    monkeypatch.setattr(platform, "home", lambda: tmp_path)
    monkeypatch.setattr(files, "home", lambda: tmp_path)
    (tmp_path / "Documents").mkdir()
    files.set_data_dir(tmp_path / ".orion")
    return tmp_path


MD = """# Quarterly report

Revenue grew **12%** thanks to *new* customers. See [site](https://example.com).

- First point
- Second point

1. Step one
2. Step two

| Region | Sales |
|---|---|
| North | 120 |
| South | 95 |

```
print("hi")
```
"""


def test_markdown_parser():
    kinds = [b[0] for b in office.md_blocks(MD)]
    assert kinds == ["h", "p", "ul", "ol", "table", "code"]
    assert office.md_blocks(MD)[4][1] == [["Region", "Sales"], ["North", "120"], ["South", "95"]]
    assert office.md_runs("a **b** *c* `d`") == [("a ", set()), ("b", {"b"}), (" ", set()), ("c", {"i"}), (" ", set()),
                                                  ("d", {"code"})]


def test_word_create_read_edit_and_undo(home):
    msg = office.create_document("documents/Report.docx", MD, title="Q3")
    assert "Word document saved" in msg
    p = home / "Documents" / "Report.docx"
    text = files.read_text_any(p)
    assert "Revenue grew 12%" in text and "North | 120" in text and "Q3" in text
    with pytest.raises(ToolError):
        office.create_document("documents/Report.docx", "again")  # no silent overwrite
    office.edit_document("documents/Report.docx", find="North", replace="Nord", append="## Appendix\n\nMore.")
    text = files.read_text_any(p)
    assert "Nord | 120" in text and "Appendix" in text
    files.undo_last_change()  # the edit is undoable: the original comes back from the backup
    assert "North | 120" in files.read_text_any(p) and "Appendix" not in files.read_text_any(p)
    with pytest.raises(ToolError):
        office.edit_document("documents/Report.docx", find="does-not-exist")


def test_pdf_create(home):
    msg = office.create_document("documents/Report.pdf", MD + "\n\nUnicode: café — ✓", title="Q3")
    assert "PDF (1 page)" in msg
    text = files.read_text_any(home / "Documents" / "Report.pdf")
    assert "Quarterly report" in text and "Revenue" in text


def test_spreadsheet_formulas_chart_update_read(home):
    rows = [["Item", "Cost"], ["Rent", "1200"], ["Food", "350.5"], ["Total", "=SUM(B2:B3)"]]
    assert "with bar chart" in office.create_spreadsheet("documents/Budget.xlsx", rows, chart="bar")
    from openpyxl import load_workbook

    ws = load_workbook(home / "Documents" / "Budget.xlsx").active
    assert ws["B2"].value == 1200 and ws["B3"].value == 350.5 and ws["B4"].value == "=SUM(B2:B3)"
    assert ws.freeze_panes == "A2" and len(ws._charts) == 1
    office.update_spreadsheet("documents/Budget.xlsx", cells=[["B2", "1300"]], append_rows=[["Fun", "80"]])
    office.update_spreadsheet("documents/Budget.xlsx", sheet="Notes", new_sheet=True, append_rows=[["Note"], ["hi"]])
    text = office.read_spreadsheet("documents/Budget.xlsx")
    assert "2: Rent | 1300" in text and "=SUM(B2:B3)" in text and "Fun | 80" in text and "Sheet: Notes" in text
    with pytest.raises(ToolError):
        office.update_spreadsheet("documents/Budget.xlsx", cells=[["not a cell", "1"]])
    assert "CSV (2 rows)" in office.create_spreadsheet("documents/x.csv", [["a", "b"], ["1", "2"]])
    assert "2: 1 | 2" in office.read_spreadsheet("documents/x.csv")


def test_presentation(home):
    from PIL import Image

    Image.new("RGB", (400, 300), "white").save(home / "pic.png")
    outline = "# Why ORION\n- Fast\n  - Really fast\n> Mention the free tier\n\n# Demo\n- Live\n![](pic.png)"
    assert "3 slides" in office.create_presentation("documents/Pitch.pptx", outline, title="ORION", subtitle="2026")
    text = files.read_text_any(home / "Documents" / "Pitch.pptx")
    assert "Why ORION" in text and "Really fast" in text and "Notes: Mention the free tier" in text


def test_pdf_tools(home):
    from PIL import Image

    office.create_document("documents/a.pdf", "# A\n\none")
    office.create_document("documents/b.pdf", "# B\n\n" + "\n\n".join(f"Para {i} " * 40 for i in range(60)))
    Image.new("RGB", (300, 200), "red").save(home / "scan.jpg")
    msg = office.combine_pdfs(["documents/a.pdf", "documents/b.pdf", "scan.jpg"], "documents/all.pdf")
    from pypdf import PdfReader

    total = len(PdfReader(str(home / "Documents" / "all.pdf")).pages)
    assert f"{total} pages from 3 files" in msg and total >= 4
    office.extract_pdf_pages("documents/all.pdf", "end, 1", output="documents/two.pdf", rotate=90)
    two = PdfReader(str(home / "Documents" / "two.pdf"))
    assert len(two.pages) == 2 and two.pages[0].rotation == 90
    with pytest.raises(ToolError):
        office.extract_pdf_pages("documents/all.pdf", "99")


def test_pdf_form_fill(home):
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, NumberObject, TextStringObject

    w = PdfWriter()
    page = w.add_blank_page(300, 300)
    field = DictionaryObject({
        NameObject("/FT"): NameObject("/Tx"), NameObject("/T"): TextStringObject("Name"),
        NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
        NameObject("/Rect"): ArrayObject([NumberObject(10), NumberObject(10), NumberObject(200), NumberObject(40)]),
    })
    ref = w._add_object(field)
    page[NameObject("/Annots")] = ArrayObject([ref])
    w._root_object[NameObject("/AcroForm")] = DictionaryObject({NameObject("/Fields"): ArrayObject([ref])})
    with open(home / "form.pdf", "wb") as fh:
        w.write(fh)
    assert office.pdf_form_fields("form.pdf")[0]["name"] == "Name"
    assert "Filled form (1 fields)" in office.fill_pdf_form("form.pdf", [["Name", "Lluc"]])
    assert office.pdf_form_fields("form (filled).pdf")[0]["value"] == "Lluc"
    with pytest.raises(ToolError):
        office.fill_pdf_form("form.pdf", [["Nope", "x"]])


def test_images_batch(home):
    from PIL import Image

    d = home / "Pictures"
    d.mkdir()
    for i in range(3):
        Image.new("RGBA", (800, 600), (i * 50, 0, 0, 255)).save(d / f"p{i}.png")
    msg = media.edit_images("pictures/*.png", output_folder="pictures/small", width=400, convert_to="jpg", quality=70)
    assert "Saved 3 image(s)" in msg
    with Image.open(d / "small" / "p0.jpg") as im:
        assert im.size == (400, 300) and im.format == "JPEG"
    media.edit_images("pictures/p1.png", rotate=90, grayscale=True)
    with Image.open(d / "p1-edited.png") as im:
        assert im.size == (600, 800) and im.mode == "L"


def test_zip_roundtrip_and_zip_slip(home):
    src = home / "proj"
    (src / "sub").mkdir(parents=True)
    (src / "a.txt").write_text("hello")
    (src / "sub" / "b.txt").write_text("world")
    assert "Zipped 2 file(s)" in media.zip_files(["proj"], "backup.zip")
    assert "Extracted 2 file(s)" in media.extract_archive("backup.zip", "restored")
    assert (home / "restored" / "proj" / "sub" / "b.txt").read_text() == "world"
    with zipfile.ZipFile(home / "evil.zip", "w") as zf:
        zf.writestr("../../escape.txt", "x")
    with pytest.raises(ToolError):
        media.extract_archive("evil.zip")


def test_search_inside_files(home):
    office.create_document("documents/Contract.docx", "The **renewal date** is 3 March.")
    (home / "Documents" / "notes.txt").write_text("nothing here")
    hits = media.search_inside_files("renewal date")
    assert len(hits) == 1 and hits[0]["file"].endswith("Contract.docx") and "3 March" in hits[0]["snippet"]
    assert "No file" in media.search_inside_files("zebra")


def test_writes_stay_in_user_folders(home):
    with pytest.raises(ToolError):
        office.create_spreadsheet("/etc/orion-test.xlsx", [["a"]])


def test_documents_tools_are_a_toolkit():
    from orion.tools.registry import decorated_tools

    kit = {t.name for t in decorated_tools() if t.kit == "documents"}
    assert {"create_document", "create_spreadsheet", "create_presentation", "combine_pdfs", "edit_images",
            "zip_files", "download_file", "search_inside_files"} <= kit
