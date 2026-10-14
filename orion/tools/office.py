"""Documents toolkit: real Office files and PDFs.

Word (.docx), Excel (.xlsx/.csv) and PowerPoint (.pptx) are written natively, so they open in Microsoft
Office, LibreOffice, Google Docs and Pages, and no Office install is needed. PDFs are produced directly
(with a Unicode font from Windows). Content is passed as Markdown, which the model writes reliably,
and converted to proper headings, lists, tables and formatting.

Every write goes through the same guard rails as write_file: user folders only, existing files are
backed up first, and changes are journaled so "undo that" works.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Literal

from .files import _backup, _check_writable, _human, _record
from .platform import IS_WIN, resolve_path
from .registry import ToolError, tool

KIT = "documents"


# ── Markdown → blocks ──────────────────────────────────────────────────────
def md_blocks(md: str) -> list[tuple]:
    """Parse the Markdown subset the model uses into blocks:
    ("h", level, text) ("p", text) ("ul", [items]) ("ol", [items]) ("table", rows) ("code", text) ("hr",)."""
    lines = md.replace("\r\n", "\n").split("\n")
    blocks: list[tuple] = []
    i = 0
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code)))
        elif m := re.match(r"^(#{1,6})\s+(.*)$", s):
            flush()
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
        elif re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush()
            blocks.append(("hr",))
        elif s.startswith("|") and s.count("|") >= 2:
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        elif re.match(r"^[-*+•]\s+", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*[-*+•]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*+•]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ul", items))
            continue
        elif re.match(r"^\d+[.)]\s+", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ol", items))
            continue
        elif not s:
            flush()
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def md_runs(text: str) -> list[tuple[str, set[str]]]:
    """Inline Markdown → [(text, {"b","i","code"})]; links become "text (url)"."""
    out: list[tuple[str, set[str]]] = []
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith(("**", "__")) and len(piece) > 4:
            out.append((piece[2:-2], {"b"}))
        elif piece.startswith("`"):
            out.append((piece[1:-1], {"code"}))
        elif piece.startswith("[") and (m := re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece)):
            out.append((f"{m.group(1)} ({m.group(2)})", set()))
        elif piece[0] in "*_" and piece[-1] == piece[0] and len(piece) > 2:
            out.append((piece[1:-1], {"i"}))
        else:
            out.append((piece, set()))
    return out


def plain(text: str) -> str:
    return "".join(t for t, _ in md_runs(text))


# ── helpers ────────────────────────────────────────────────────────────────
def _target(path: str, exts: tuple[str, ...], overwrite: bool = True) -> tuple[Path, str | None, bool]:
    p = resolve_path(path)
    if p.suffix.lower() not in exts:
        p = p.with_suffix(exts[0])
    _check_writable(p)
    if p.exists() and not overwrite:
        raise ToolError(f"{p} already exists. Pick another name or allow overwriting.")
    p.parent.mkdir(parents=True, exist_ok=True)
    existed = p.exists()
    return p, (_backup(p) if existed else None), existed


def _done(p: Path, backup: str | None, existed: bool, what: str) -> str:
    if not p.exists() or p.stat().st_size == 0:
        raise ToolError(f"{p.name} could not be verified after saving.")
    _record({"op": "write", "path": str(p), "backup": backup} if existed else {"op": "create", "path": str(p)})
    return f"{what} saved: {p} ({_human(p.stat().st_size)})."


def _num(v: str | float | int | None):
    """Spreadsheet cell value: formulas stay text, numbers become numbers."""
    if v is None or isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s.startswith("="):
        return s
    t = s.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", s) else s
    for conv in (int, float):
        try:
            return conv(t)
        except ValueError:
            pass
    if re.fullmatch(r"-?\d+(\.\d+)?%", s):
        return float(s[:-1]) / 100
    return s


# ── Word ───────────────────────────────────────────────────────────────────
def _docx_runs(par, text: str) -> None:
    for chunk, style in md_runs(text):
        run = par.add_run(chunk)
        run.bold = "b" in style or None
        run.italic = "i" in style or None
        if "code" in style:
            run.font.name = "Consolas"


def _docx_blocks(doc, blocks: list[tuple]) -> None:
    from docx.shared import Pt

    for b in blocks:
        kind = b[0]
        if kind == "h":
            _docx_runs(doc.add_heading(level=min(b[1], 4)), b[2])
        elif kind == "p":
            _docx_runs(doc.add_paragraph(), b[1])
        elif kind in ("ul", "ol"):
            for item in b[1]:
                _docx_runs(doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), item)
        elif kind == "table" and b[1]:
            rows = b[1]
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _docx_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif kind == "code":
            par = doc.add_paragraph()
            run = par.add_run(b[1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif kind == "hr":
            doc.add_paragraph("—" * 30)


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Create document {a.get('path')}")
def create_document(path: str, content: str, title: str = "", overwrite: bool = False) -> str:
    """Create a formatted document: Word (.docx) by default, or PDF if the path ends in .pdf.
    Write the content in Markdown: # headings, paragraphs, - bullets, 1. numbered lists, | tables |,
    **bold**, *italic*, `code` and ``` code blocks. Use for letters, reports, CVs, essays, notes.

    Args:
        path: Where to save, e.g. "documents/Cover letter.docx" or "desktop/report.pdf".
        content: The document body in Markdown.
        title: Optional title shown at the top.
        overwrite: Replace the file if it already exists.
    """
    if resolve_path(path).suffix.lower() == ".pdf":
        return create_pdf(path, content, title, overwrite)
    import docx

    p, backup, existed = _target(path, (".docx",), overwrite)
    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)
        doc.core_properties.title = title
    _docx_blocks(doc, md_blocks(content))
    doc.save(str(p))
    docx.Document(str(p))  # read-back check: the file opens
    return _done(p, backup, existed, "Word document")


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Edit {a.get('path')}")
def edit_document(path: str, find: str = "", replace: str = "", append: str = "") -> str:
    """Edit an existing Word document: replace text everywhere it appears and/or append Markdown content
    at the end. The original is backed up so the change can be undone.

    Args:
        path: The .docx file.
        find: Exact text to find (leave empty to only append).
        replace: Replacement text.
        append: Markdown to add at the end (same syntax as create_document).
    """
    import docx

    p = resolve_path(path)
    if p.suffix.lower() != ".docx" or not p.is_file():
        raise ToolError(f"Not a Word document: {p}")
    _check_writable(p)
    backup = _backup(p)
    doc = docx.Document(str(p))
    count = 0
    if find:
        pars = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    pars.extend(cell.paragraphs)
        for par in pars:
            if find not in par.text:
                continue
            n = par.text.count(find)
            if any(find in r.text for r in par.runs):  # keep formatting when the text sits inside one run
                for r in par.runs:
                    if find in r.text:
                        r.text = r.text.replace(find, replace)
            else:
                full = par.text.replace(find, replace)
                for r in par.runs[1:]:
                    r.text = ""
                par.runs[0].text = full
            count += n
        if not count:
            raise ToolError(f"“{find}” wasn't found in {p.name}.")
    if append:
        _docx_blocks(doc, md_blocks(append))
    doc.save(str(p))
    docx.Document(str(p))
    _record({"op": "write", "path": str(p), "backup": backup})
    parts = ([f"replaced {count} occurrence(s)"] if find else []) + (["appended content"] if append else [])
    return f"Updated {p.name}: {', '.join(parts) or 'no changes'}."


# ── PDF ────────────────────────────────────────────────────────────────────
def _font_files() -> dict[str, str] | None:
    """A Unicode TTF family for PDFs (regular required; missing styles fall back to regular/bold)."""
    import os

    families: list[tuple[Path, tuple[str, str, str, str]]] = []
    if IS_WIN:
        fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
        families += [(fonts, ("segoeui.ttf", "segoeuib.ttf", "segoeuii.ttf", "segoeuiz.ttf")),
                     (fonts, ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"))]
    families += [(Path("/usr/share/fonts/truetype/liberation"),
                  ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf",
                   "LiberationSans-BoldItalic.ttf")),
                 (Path("/usr/share/fonts/truetype/dejavu"),
                  ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf")),
                 (Path("/System/Library/Fonts/Supplemental"), ("Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf",
                                                              "Arial Bold Italic.ttf"))]
    for folder, names in families:
        paths = [folder / n for n in names]
        if not paths[0].exists():
            continue
        reg = str(paths[0])
        bold = str(paths[1]) if paths[1].exists() else reg
        ital = str(paths[2]) if paths[2].exists() else reg
        both = str(paths[3]) if paths[3].exists() else bold
        return {"": reg, "B": bold, "I": ital, "BI": both}
    return None


def _symbol_fonts() -> list[str]:
    """Fonts with symbols/arrows/ticks for glyphs the body font lacks."""
    import os

    cands = []
    if IS_WIN:
        fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
        cands += [fonts / "seguisym.ttf", fonts / "arialuni.ttf"]
    cands += [Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    return [str(p) for p in cands if p.exists()]


def _html_inline(text: str) -> str:
    import html

    out = []
    for chunk, style in md_runs(text):
        s = html.escape(chunk)
        if "code" in style:
            s = f"<code>{s}</code>"
        if "b" in style:
            s = f"<b>{s}</b>"
        if "i" in style:
            s = f"<i>{s}</i>"
        out.append(s)
    return "".join(out)


def md_to_html(blocks: list[tuple]) -> str:
    import html

    out = []
    for b in blocks:
        kind = b[0]
        if kind == "h":
            out.append(f"<h{min(b[1], 4)}>{_html_inline(b[2])}</h{min(b[1], 4)}>")
        elif kind == "p":
            out.append(f"<p>{_html_inline(b[1])}</p>")
        elif kind in ("ul", "ol"):
            out.append(f"<{kind}>" + "".join(f"<li>{_html_inline(i)}</li>" for i in b[1]) + f"</{kind}>")
        elif kind == "table" and b[1]:
            width = max(len(r) for r in b[1])
            rows = []
            for r, row in enumerate(b[1]):
                tag = "th" if r == 0 else "td"
                rows.append("<tr>" + "".join(f"<{tag}>{_html_inline(row[c] if c < len(row) else '')}</{tag}>"
                                             for c in range(width)) + "</tr>")
            out.append('<table width="100%">' + "".join(rows) + "</table>")
        elif kind == "code":
            out.append(f"<pre>{html.escape(b[1])}</pre>")
        elif kind == "hr":
            out.append("<hr>")
    return "\n".join(out)


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Create PDF {a.get('path')}")
def create_pdf(path: str, content: str, title: str = "", overwrite: bool = False) -> str:
    """Create a PDF document from Markdown (headings, paragraphs, lists, tables, bold/italic, code).

    Args:
        path: Where to save, e.g. "documents/Invoice 042.pdf".
        content: The document body in Markdown.
        title: Optional title shown at the top.
        overwrite: Replace the file if it already exists.
    """
    from fpdf import FPDF

    p, backup, existed = _target(path, (".pdf",), overwrite)
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, 18)
    fonts = _font_files()
    family = "helvetica"
    if fonts:
        for style, f in fonts.items():
            pdf.add_font("body", style, f)
        family = "body"
        fallbacks = []
        for n, f in enumerate(_symbol_fonts()):
            pdf.add_font(f"sym{n}", "", f)
            fallbacks.append(f"sym{n}")
        if fallbacks:
            pdf.set_fallback_fonts(fallbacks, exact_match=False)
    pdf.add_page()
    if title:
        pdf.set_title(title)
    body = (f"<h1>{_html_inline(title)}</h1>" if title else "") + md_to_html(md_blocks(content))
    if not fonts:  # core fonts are Latin-1 only
        body = body.encode("latin-1", "replace").decode("latin-1")
    pdf.set_font(family, size=11)
    from fpdf.fonts import TextStyle

    ink = (20, 20, 20)
    heading = {f"h{n}": TextStyle(font_family=family, font_style="B", font_size_pt=size, color=ink, t_margin=5, b_margin=2)
               for n, size in ((1, 20), (2, 16), (3, 13), (4, 12))}
    pdf.write_html(body, font_family=family, warn_on_tags_not_matching=False, li_prefix_color=(90, 90, 90),
                   tag_styles={**heading, "pre": TextStyle(font_family="courier", font_size_pt=9.5, color=ink),
                               "code": TextStyle(font_family="courier", font_size_pt=10, color=ink)})
    pdf.output(str(p))
    from pypdf import PdfReader

    pages = len(PdfReader(str(p)).pages)
    return _done(p, backup, existed, f"PDF ({pages} page{'s' if pages != 1 else ''})")


# ── Excel ──────────────────────────────────────────────────────────────────
def _style_sheet(ws, header: bool) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if header and ws.max_row >= 1:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="262626")
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None and not str(c.value).startswith("=")),
                    default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(9, width + 2))


def _add_chart(ws, kind: str, title: str) -> None:
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    if ws.max_row < 2 or ws.max_column < 2:
        raise ToolError("A chart needs a header row, a label column and at least one numeric column.")
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    chart.title = title or None
    last_col = 2 if kind == "pie" else ws.max_column
    data = Reference(ws, min_col=2, max_col=last_col, min_row=1, max_row=ws.max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width, chart.height = 18, 9
    from openpyxl.utils import get_column_letter

    ws.add_chart(chart, f"{get_column_letter(ws.max_column + 2)}2")


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Create spreadsheet {a.get('path')}")
def create_spreadsheet(path: str, rows: list[list[str]], sheet_name: str = "Sheet1", header: bool = True,
                       chart: Literal["none", "bar", "line", "pie"] = "none", chart_title: str = "",
                       overwrite: bool = False) -> str:
    """Create an Excel workbook (.xlsx) or CSV. Numbers become real numbers and cells starting with "="
    are live formulas (e.g. "=SUM(B2:B9)"). The header row is styled, frozen and filterable, columns are
    sized, and an optional chart plots column A (labels) against the other columns.

    Args:
        path: Where to save, e.g. "documents/Budget 2026.xlsx" (use .csv for plain CSV).
        rows: The table, row by row; the first row is the header.
        sheet_name: Name of the worksheet.
        header: Whether the first row is a header.
        chart: Add a bar, line or pie chart of the data.
        chart_title: Title for the chart.
        overwrite: Replace the file if it already exists.
    """
    if not rows:
        raise ToolError("No rows given.")
    if resolve_path(path).suffix.lower() == ".csv":
        p, backup, existed = _target(path, (".csv",), overwrite)
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows(rows)
        return _done(p, backup, existed, f"CSV ({len(rows)} rows)")
    from openpyxl import Workbook, load_workbook

    p, backup, existed = _target(path, (".xlsx",), overwrite)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "Sheet1"
    for row in rows:
        ws.append([_num(v) for v in row])
    _style_sheet(ws, header)
    if chart != "none":
        _add_chart(ws, chart, chart_title)
    wb.save(str(p))
    load_workbook(str(p))  # read-back check
    return _done(p, backup, existed, f"Spreadsheet ({len(rows)} rows × {max(len(r) for r in rows)} columns"
                                      f"{', with ' + chart + ' chart' if chart != 'none' else ''})")


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Update spreadsheet {a.get('path')}")
def update_spreadsheet(path: str, sheet: str = "", cells: list[list[str]] | None = None,
                       append_rows: list[list[str]] | None = None, new_sheet: bool = False) -> str:
    """Change an existing Excel workbook: set individual cells, append rows, or add a new sheet.
    The original is backed up so the change can be undone.

    Args:
        path: The .xlsx file.
        sheet: Worksheet name (empty = the first sheet). With new_sheet, the name of the sheet to create.
        cells: Pairs of [cell, value], e.g. [["B2", "450"], ["B10", "=SUM(B2:B9)"]].
        append_rows: Rows to add after the last used row.
        new_sheet: Create the sheet named in `sheet` (append_rows then fills it).
    """
    from openpyxl import load_workbook

    p = resolve_path(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm") or not p.is_file():
        raise ToolError(f"Not an Excel workbook: {p}")
    _check_writable(p)
    backup = _backup(p)
    wb = load_workbook(str(p))
    if new_sheet:
        if not sheet:
            raise ToolError("Give the new sheet a name.")
        ws = wb.create_sheet(sheet[:31])
    elif sheet:
        if sheet not in wb.sheetnames:
            raise ToolError(f"No sheet “{sheet}”. Sheets: {', '.join(wb.sheetnames)}")
        ws = wb[sheet]
    else:
        ws = wb.worksheets[0]
    changed = 0
    for pair in cells or []:
        if len(pair) < 2 or not re.fullmatch(r"[A-Za-z]{1,3}\d{1,7}", str(pair[0]).strip()):
            raise ToolError(f"Bad cell entry {pair!r}; use [\"B2\", \"value\"].")
        ws[str(pair[0]).strip().upper()] = _num(pair[1])
        changed += 1
    for row in append_rows or []:
        ws.append([_num(v) for v in row])
    if new_sheet and append_rows:
        _style_sheet(ws, True)
    wb.save(str(p))
    load_workbook(str(p))
    _record({"op": "write", "path": str(p), "backup": backup})
    return (f"Updated {p.name} [{ws.title}]: {changed} cell(s) set, {len(append_rows or [])} row(s) added"
            f"{' (new sheet)' if new_sheet else ''}.")


def read_spreadsheet_text(p: Path, sheet: str = "", max_rows: int = 200) -> str:
    from openpyxl import load_workbook

    values = load_workbook(str(p), data_only=True, read_only=True)
    formulas = load_workbook(str(p), read_only=True)
    names = [sheet] if sheet else values.sheetnames
    out = []
    for name in names:
        if name not in values.sheetnames:
            raise ToolError(f"No sheet “{name}”. Sheets: {', '.join(values.sheetnames)}")
        wv, wf = values[name], formulas[name]
        out.append(f"## Sheet: {name}")
        for r, (rv, rf) in enumerate(zip(wv.iter_rows(values_only=True), wf.iter_rows(values_only=True), strict=False)):
            if r >= max_rows:
                out.append(f"…(more rows; {wv.max_row} total)")
                break
            cells = []
            for v, f in zip(rv, rf, strict=False):
                if v is None and isinstance(f, str) and f.startswith("="):
                    cells.append(f)  # formula never calculated (file not opened in Excel yet)
                else:
                    cells.append("" if v is None else str(v))
            if any(cells):
                out.append(f"{r + 1}: " + " | ".join(cells))
    return "\n".join(out)


@tool(category="documents", kit=KIT)
def read_spreadsheet(path: str, sheet: str = "", max_rows: int = 200) -> str:
    """Read an Excel workbook (.xlsx) or CSV with row numbers, so you can analyse it or target cells to update.

    Args:
        path: The file.
        sheet: A specific sheet (empty = all sheets).
        max_rows: Row limit per sheet.
    """
    p = resolve_path(path)
    if not p.is_file():
        raise ToolError(f"File not found: {p}")
    if p.suffix.lower() == ".csv":
        with open(p, newline="", encoding="utf-8-sig", errors="replace") as f:
            rows = list(csv.reader(f))
        lines = [f"{i + 1}: " + " | ".join(r) for i, r in enumerate(rows[:max_rows])]
        return "\n".join(lines + ([f"…({len(rows)} rows total)"] if len(rows) > max_rows else []))
    return read_spreadsheet_text(p, sheet, max_rows)


# ── PowerPoint ─────────────────────────────────────────────────────────────
def outline_slides(outline: str) -> list[dict]:
    """'# Slide title' starts a slide; '- ' bullets (indent for sub-bullets); '> ' speaker notes;
    '![](path)' adds an image; other lines become body text."""
    slides: list[dict] = []
    cur: dict | None = None
    for raw in outline.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        s = line.strip()
        if not s:
            continue
        if m := re.match(r"^#{1,3}\s+(.*)$", s):
            cur = {"title": m.group(1).strip(), "bullets": [], "notes": [], "image": None}
            slides.append(cur)
            continue
        if cur is None:
            cur = {"title": "", "bullets": [], "notes": [], "image": None}
            slides.append(cur)
        if m := re.match(r"^!\[[^\]]*\]\(([^)]+)\)$", s):
            cur["image"] = m.group(1)
        elif s.startswith(">"):
            cur["notes"].append(s.lstrip("> ").strip())
        elif m := re.match(r"^(\s*)[-*•]\s+(.*)$", line):
            cur["bullets"].append((min(len(m.group(1).replace("\t", "  ")) // 2, 3), m.group(2).strip()))
        else:
            cur["bullets"].append((0, s))
    return slides


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Create presentation {a.get('path')}")
def create_presentation(path: str, outline: str, title: str = "", subtitle: str = "",
                        theme: Literal["dark", "light"] = "dark", overwrite: bool = False) -> str:
    """Create a PowerPoint deck (.pptx, 16:9) from an outline. Each slide starts with "# Slide title",
    followed by "- bullet" lines (indent two spaces for sub-bullets), optional "> speaker notes" and an
    optional "![](path/to/image.png)". Keep 3-6 short bullets per slide.

    Args:
        path: Where to save, e.g. "documents/Pitch.pptx".
        outline: The slides in the format above.
        title: Title slide heading (adds a title slide first).
        subtitle: Title slide subheading.
        theme: Dark or light colours.
        overwrite: Replace the file if it already exists.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    slides = outline_slides(outline)
    if not slides and not title:
        raise ToolError("The outline has no slides. Start each slide with '# Title'.")
    p, backup, existed = _target(path, (".pptx",), overwrite)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    bg, fg, muted = ((RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0xF2, 0xF2, 0xF2), RGBColor(0xA6, 0xA6, 0xA6))
                     if theme == "dark" else
                     (RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x14, 0x14, 0x14), RGBColor(0x59, 0x59, 0x59)))

    def paint(slide) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = bg
        for shape in slide.placeholders:
            if shape.has_text_frame:
                for par in shape.text_frame.paragraphs:
                    for run in par.runs:
                        run.font.color.rgb = fg

    def fit(shape, left, top, width, height) -> None:
        shape.left, shape.top, shape.width, shape.height = Emu(left), Emu(top), Emu(width), Emu(height)

    W, H = prs.slide_width, prs.slide_height
    if title:
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = title
        fit(s.shapes.title, Inches(0.8), Inches(2.4), W - Inches(1.6), Inches(1.5))
        sub = s.placeholders[1]
        sub.text = subtitle
        fit(sub, Inches(0.8), Inches(4.0), W - Inches(1.6), Inches(1.0))
        paint(s)
        for run in (r for par in sub.text_frame.paragraphs for r in par.runs):
            run.font.color.rgb = muted
    for spec in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = plain(spec["title"])
        fit(s.shapes.title, Inches(0.7), Inches(0.4), W - Inches(1.4), Inches(1.1))
        body = s.placeholders[1]
        img = spec["image"]
        text_w = W - Inches(1.4)
        if img:
            ip = resolve_path(img)
            if not ip.is_file():
                raise ToolError(f"Image not found for slide “{spec['title']}”: {ip}")
            text_w = int(W * 0.5)
            pic = s.shapes.add_picture(str(ip), Inches(0.7) + text_w + Inches(0.3), Inches(1.7))
            max_w, max_h = W - text_w - Inches(1.7), H - Inches(2.3)
            scale = min(max_w / pic.width, max_h / pic.height, 1)
            pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        fit(body, Inches(0.7), Inches(1.6), text_w, H - Inches(2.2))
        tf = body.text_frame
        tf.clear()
        for n, (level, text) in enumerate(spec["bullets"] or [(0, "")]):
            par = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            par.level = level
            for chunk, style in md_runs(text):
                run = par.add_run()
                run.text = chunk
                run.font.bold = "b" in style or None
                run.font.italic = "i" in style or None
                run.font.size = Pt(24 if level == 0 else 20)
        if spec["notes"]:
            s.notes_slide.notes_text_frame.text = "\n".join(spec["notes"])
        paint(s)
    prs.save(str(p))
    count = len(Presentation(str(p)).slides)
    return _done(p, backup, existed, f"Presentation ({count} slides)")


def read_presentation_text(p: Path) -> str:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        out.append(f"## Slide {i}\n" + "\n".join(texts))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
    return "\n\n".join(out)


# ── PDF utilities ──────────────────────────────────────────────────────────
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}


def _page_list(spec: str, total: int) -> list[int]:
    pages: list[int] = []
    for part in re.split(r"[,\s]+", spec.strip()):
        if not part:
            continue
        m = re.fullmatch(r"(\d+|end)?\s*-\s*(\d+|end)?", part)
        if m:
            a = 1 if not m.group(1) else (total if m.group(1) == "end" else int(m.group(1)))
            b = total if not m.group(2) or m.group(2) == "end" else int(m.group(2))
            pages.extend(range(a, b + 1) if a <= b else range(a, b - 1, -1))
        elif part.isdigit() or part == "end":
            pages.append(total if part == "end" else int(part))
        else:
            raise ToolError(f"Can't read the page list “{spec}”. Use e.g. “1-3, 5, 8-end”.")
    bad = [n for n in pages if n < 1 or n > total]
    if bad:
        raise ToolError(f"Page {bad[0]} doesn't exist; the PDF has {total} pages.")
    return pages


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Combine into {a.get('output')}")
def combine_pdfs(files: list[str], output: str) -> str:
    """Merge PDFs and/or images (photos, scans) into one PDF, in the given order.

    Args:
        files: The PDFs and images to combine, in order.
        output: The new PDF's path.
    """
    import io

    from PIL import Image
    from pypdf import PdfReader, PdfWriter

    if len(files) < 1:
        raise ToolError("Give at least one file.")
    writer = PdfWriter()
    for f in files:
        fp = resolve_path(f)
        if not fp.is_file():
            raise ToolError(f"File not found: {fp}")
        if fp.suffix.lower() in IMAGE_EXTS:
            buf = io.BytesIO()
            with Image.open(fp) as im:
                im.convert("RGB").save(buf, "PDF", resolution=150)
            buf.seek(0)
            writer.append(PdfReader(buf))
        elif fp.suffix.lower() == ".pdf":
            writer.append(str(fp))
        else:
            raise ToolError(f"{fp.name} is not a PDF or image.")
    p, backup, existed = _target(output, (".pdf",))
    with open(p, "wb") as fh:
        writer.write(fh)
    pages = len(PdfReader(str(p)).pages)
    return _done(p, backup, existed, f"Combined PDF ({pages} pages from {len(files)} files)")


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Pages {a.get('pages')} of {a.get('path')}")
def extract_pdf_pages(path: str, pages: str, output: str = "", rotate: int = 0) -> str:
    """Make a new PDF from selected pages of a PDF (split, reorder, remove pages) and optionally rotate them.

    Args:
        path: The source PDF.
        pages: Pages to keep in order, e.g. "1-3, 5, 8-end" or "end-1" to reverse.
        output: New file path (default: "<name> (pages ...).pdf" next to the original).
        rotate: Rotate the kept pages clockwise by 90, 180 or 270 degrees.
    """
    from pypdf import PdfReader, PdfWriter

    src = resolve_path(path)
    if not src.is_file():
        raise ToolError(f"File not found: {src}")
    reader = PdfReader(str(src))
    keep = _page_list(pages, len(reader.pages))
    writer = PdfWriter()
    for n in keep:
        page = writer.add_page(reader.pages[n - 1])
        if rotate % 360:
            page.rotate(rotate % 360)
    safe = re.sub(r"[^\w,-]+", " ", pages).strip()
    p, backup, existed = _target(output or str(src.with_name(f"{src.stem} (pages {safe}).pdf")), (".pdf",))
    with open(p, "wb") as fh:
        writer.write(fh)
    return _done(p, backup, existed, f"PDF with {len(keep)} page(s)")


@tool(category="documents", kit=KIT)
def pdf_form_fields(path: str) -> list[dict] | str:
    """List the fillable fields of a PDF form (name, type, current value, choices).

    Args:
        path: The PDF form.
    """
    from pypdf import PdfReader

    p = resolve_path(path)
    if not p.is_file():
        raise ToolError(f"File not found: {p}")
    fields = PdfReader(str(p)).get_fields() or {}
    if not fields:
        return "This PDF has no fillable form fields. (To fill a flat PDF, describe it and create a new one.)"
    kinds = {"/Tx": "text", "/Btn": "checkbox/button", "/Ch": "choice", "/Sig": "signature"}
    out = []
    for name, f in fields.items():
        item = {"name": name, "type": kinds.get(str(f.get("/FT")), str(f.get("/FT"))), "value": str(f.get("/V") or "")}
        if f.get("/_States_"):
            item["choices"] = [str(s) for s in f["/_States_"]]
        elif f.get("/Opt"):
            item["choices"] = [str(o[-1] if isinstance(o, list) else o) for o in f["/Opt"]]
        out.append(item)
    return out


@tool(risk="confirm", category="documents", kit=KIT, summary=lambda a: f"Fill form {a.get('path')}")
def fill_pdf_form(path: str, values: list[list[str]], output: str = "") -> str:
    """Fill a PDF form's fields (get the names from pdf_form_fields) and save a filled copy.

    Args:
        path: The PDF form.
        values: Pairs of [field name, value]. For checkboxes use the checked state, e.g. "/Yes".
        output: Where to save (default: "<name> (filled).pdf").
    """
    from pypdf import PdfReader, PdfWriter

    src = resolve_path(path)
    if not src.is_file():
        raise ToolError(f"File not found: {src}")
    reader = PdfReader(str(src))
    fields = reader.get_fields() or {}
    data = {str(k): str(v) for k, v, *_ in values}
    unknown = [k for k in data if k not in fields]
    if unknown:
        raise ToolError(f"Unknown field(s): {', '.join(unknown)}. Use pdf_form_fields for the exact names.")
    writer = PdfWriter(clone_from=reader)
    for page in writer.pages:
        writer.update_page_form_field_values(page, data, auto_regenerate=False)
    writer.set_need_appearances_writer(True)
    p, backup, existed = _target(output or str(src.with_name(f"{src.stem} (filled).pdf")), (".pdf",))
    with open(p, "wb") as fh:
        writer.write(fh)
    filled = PdfReader(str(p)).get_fields() or {}
    missing = [k for k, v in data.items() if str((filled.get(k) or {}).get("/V", "")) != v]
    if missing:
        raise ToolError(f"Saved, but these fields didn't take the value: {', '.join(missing)}")
    return _done(p, backup, existed, f"Filled form ({len(data)} fields)")
