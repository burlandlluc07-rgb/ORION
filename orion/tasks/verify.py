"""Independent result checks: no action counts as done until its result has been looked at.

A verifier gets the tool's arguments and its result text and returns::

    {"status": "passed" | "failed" | "unverified" | "not_applicable", "evidence": [...], "reason": "..."}

``failed`` means the action reported success but the check found the result wrong or missing; the
step then ends as ``needs_attention`` instead of ``completed``. ``unverified`` means nothing could be
checked independently, which is recorded honestly rather than counted as proof.
"""

from __future__ import annotations

import re
import time
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Callable

if TYPE_CHECKING:
    from .lifecycle import Capabilities

Verifier = Callable[[dict, str], dict]

SIGNATURES = {".docx": b"PK", ".xlsx": b"PK", ".pptx": b"PK", ".zip": b"PK", ".pdf": b"%PDF", ".png": b"\x89PNG",
              ".jpg": b"\xff\xd8", ".jpeg": b"\xff\xd8", ".gif": b"GIF8", ".dwg": b"AC10"}


def passed(*evidence: str) -> dict:
    return {"status": "passed", "evidence": [e for e in evidence if e]}


def failed(reason: str, *evidence: str) -> dict:
    return {"status": "failed", "reason": reason, "evidence": [e for e in evidence if e]}


def unverified(reason: str) -> dict:
    return {"status": "unverified", "reason": reason, "evidence": []}


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[*_`#>|\-]+", " ", text or "")).strip().lower()


def snippet(content: str, words: int = 6) -> str:
    """The first few words of what was asked for, without Markdown, to look for in the saved file."""
    from ..tools.office import md_blocks, plain

    text = content
    try:  # the first block of text as the document will show it (no list numbers, bullets or table pipes)
        for block in md_blocks(content):
            if block[0] in ("h", "p") and str(block[-1]).strip():
                text = str(block[-1])
            elif block[0] in ("ul", "ol") and block[1]:
                text = str(block[1][0])
            elif block[0] == "table" and block[1] and block[1][0]:
                text = str(block[1][0][0])
            else:
                continue
            break
        text = plain(text)
    except Exception:
        pass
    return " ".join(_norm(text).split()[:words])


def path_from(result: str, args: dict, *keys: str) -> Path | None:
    m = re.search(r"(?:saved|Downloaded|into|Screenshot saved):\s*(.+?)(?:\s+\(\d|\s*$)", result or "") or \
        re.search(r"(?:Downloaded|into)\s+(.+?)\s+\(", result or "")
    if m:
        return Path(m.group(1).strip().rstrip("."))
    from ..tools.platform import resolve_path

    for k in keys:
        if args.get(k):
            return resolve_path(str(args[k]))
    return None


def check_file(p: Path | None, min_size: int = 1) -> tuple[bool, str]:
    if p is None:
        return False, "no output path was reported"
    if not p.exists():
        return False, f"{p} does not exist"
    size = p.stat().st_size
    if size < min_size:
        return False, f"{p} is empty"
    sig = SIGNATURES.get(p.suffix.lower())
    if sig:
        with open(p, "rb") as f:
            head = f.read(len(sig))
        if head != sig:
            return False, f"{p.name} is not a valid {p.suffix} file (wrong file signature)"
    return True, f"{p.name} exists ({size} bytes, format looks right)"


def document_text(p: Path) -> str:
    ext = p.suffix.lower()
    if ext == ".docx":
        import docx

        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    if ext == ".pptx":
        from pptx import Presentation

        out = []
        for slide in Presentation(str(p)).slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    out.append(shape.text_frame.text)
        return "\n".join(out)
    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(str(p), read_only=True)
        cells = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells.extend(str(v) for v in row if v is not None)
        wb.close()
        return "\n".join(cells)
    return p.read_text(encoding="utf-8", errors="replace")


def _document(content_key: str, title_key: str | None = "title", path_keys: tuple[str, ...] = ("path",),
              words: int = 6) -> Verifier:
    def check(args: dict, result: str) -> dict:
        p = path_from(result, args, *path_keys)
        ok, why = check_file(p)
        if not ok:
            return failed(why)
        try:
            text = _norm(document_text(p))  # reopen the file the way a user's app would
        except Exception as exc:
            return failed(f"{p.name} could not be reopened: {type(exc).__name__}: {exc}", why)
        wanted = [s for s in (snippet(str(args.get(content_key) or ""), words),
                              _norm(str(args.get(title_key) or "")) if title_key else "") if s]
        missing = [w for w in wanted if w not in text]
        if missing:
            return failed(f"{p.name} opened, but the expected content is missing: “{missing[0][:60]}”", why)
        return passed(why, f"reopened {p.name} and found the expected content" if wanted else f"reopened {p.name}")
    return check


def _spreadsheet(args: dict, result: str) -> dict:
    p = path_from(result, args, "path")
    ok, why = check_file(p)
    if not ok:
        return failed(why)
    try:
        text = _norm(document_text(p))
    except Exception as exc:
        return failed(f"{p.name} could not be reopened: {exc}", why)
    first = next((str(c) for r in (args.get("rows") or []) for c in r if str(c).strip() and not str(c).startswith("=")),
                 "")
    if first and _norm(first) not in text:
        return failed(f"{p.name} opened, but the cell “{first[:40]}” is missing", why)
    return passed(why, f"reopened {p.name} and found the expected cells")


def _pdf_pages(args: dict, result: str) -> dict:
    p = path_from(result, args, "output")
    ok, why = check_file(p)
    if not ok:
        return failed(why)
    try:
        from pypdf import PdfReader

        n = len(PdfReader(str(p)).pages)
    except Exception as exc:
        return failed(f"{p.name} could not be reopened: {exc}", why)
    return passed(why, f"reopened {p.name}: {n} page(s)") if n else failed(f"{p.name} has no pages", why)


def _zip(args: dict, result: str) -> dict:
    p = path_from(result, args, "output")
    ok, why = check_file(p)
    if not ok:
        return failed(why)
    try:
        with zipfile.ZipFile(p) as z:
            bad = z.testzip()
            n = len(z.namelist())
    except zipfile.BadZipFile as exc:
        return failed(f"{p.name} is not a readable zip: {exc}", why)
    return failed(f"{bad} is corrupt in {p.name}", why) if bad else passed(why, f"{n} file(s) inside, all readable")


def _download(args: dict, result: str) -> dict:
    p = path_from(result, args)
    ok, why = check_file(p)
    return passed(why) if ok else failed(why)


def _screenshot(args: dict, result: str) -> dict:
    p = path_from(result, args)
    ok, why = check_file(p, min_size=100)
    return passed(why) if ok else failed(why)


def _open_app(args: dict, result: str) -> dict:
    """A launch is not proof: look for the app's process (a short wait for slow starters)."""
    import psutil

    want = re.sub(r"[^a-z0-9]", "", str(args.get("name", "")).lower())
    m = re.search(r"Opened (.+?)\.$|Started (.+?)\.$", result or "")
    shown = re.sub(r"[^a-z0-9]", "", (m.group(1) or m.group(2) or "").lower()) if m else ""
    keys = {k for k in (want, shown, want[:6]) if len(k) >= 3}
    if not keys:
        return unverified("no app name to look for")
    deadline = time.time() + 5
    while True:
        for proc in psutil.process_iter(["name"]):
            pname = re.sub(r"[^a-z0-9]", "", (proc.info["name"] or "").lower().removesuffix(".exe"))
            if any(k in pname or (len(pname) >= 4 and pname in k) for k in keys):
                return passed(f"process {proc.info['name']} is running")
        if time.time() > deadline:
            # Store apps and some launchers run under a host process with another name: not a definite failure.
            return unverified(f"launched, but no process matching “{args.get('name')}” was found")
        time.sleep(0.5)


def _calendar_event(args: dict, result: str) -> dict:
    return unverified("ORION opened a pre-filled event; it can't confirm you pressed Save")


def _self_checked(how: str) -> Verifier:
    return lambda args, result: passed(f"the tool checked its own result: {how}")


VERIFIERS: dict[str, Verifier] = {
    "create_document": _document("content"),
    "create_pdf": _document("content"),
    "create_presentation": _document("outline", words=3),
    "create_spreadsheet": _spreadsheet,
    "combine_pdfs": _pdf_pages,
    "extract_pdf_pages": _pdf_pages,
    "fill_pdf_form": _pdf_pages,
    "zip_files": _zip,
    "download_file": _download,
    "take_screenshot": _screenshot,
    "open_app": _open_app,
    "calendar_add_event": _calendar_event,
}


def verifier_for(name: str, caps: "Capabilities") -> Verifier | None:
    from .routing import SELF_CHECKED

    if name in VERIFIERS:
        return VERIFIERS[name]
    if name in SELF_CHECKED:
        return _self_checked(SELF_CHECKED[name])
    return None
